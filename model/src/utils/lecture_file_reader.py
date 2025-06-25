import pdfplumber
from docx import Document
import openpyxl
import pdfplumber
import fitz  # PyMuPDF

def read_pdf(file_path):
    text = []

    try:
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text.append(page_text)
                    else:
                        # Fallback if pdfplumber couldn't extract anything
                        raise ValueError("Empty text, try fallback")
                except Exception as e:
                    print(f"⚠️ pdfplumber failed on page {i+1}: {e}")
                    # Fallback to PyMuPDF for this page
                    try:
                        fallback_page = fitz.open(file_path)[i]
                        text.append(fallback_page.get_text())
                        print(f"✅ Recovered page {i+1} with PyMuPDF")
                    except Exception as fallback_e:
                        print(f"❌ Fallback also failed for page {i+1}: {fallback_e}")
    except Exception as outer_e:
        print(f"❌ Failed to open PDF with pdfplumber: {outer_e}")
        # Try full fallback
        try:
            doc = fitz.open(file_path)
            text = [page.get_text() for page in doc]
            print(f"✅ Entire file extracted with PyMuPDF")
        except Exception as final_e:
            print(f"❌ Completely failed to read {file_path}: {final_e}")
            return ""

    return '\n'.join(text)


def read_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

# def read_pdf(file_path):
#     with pdfplumber.open(file_path) as pdf:
#         return '\n'.join(page.extract_text() for page in pdf.pages if page.extract_text())

def read_docx(file_path):
    doc = Document(file_path)
    return '\n'.join(para.text for para in doc.paragraphs)

def read_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    content = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = ' | '.join(str(cell) for cell in row if cell is not None)
            if line:
                content.append(line)
    return '\n'.join(content)

def read_file(file_path):
    if file_path.endswith('.txt'):
        return read_txt(file_path)
    elif file_path.endswith('.pdf'):
        return read_pdf(file_path)
    elif file_path.endswith('.docx'):
        return read_docx(file_path)
    elif file_path.endswith('.xlsx') or file_path.endswith('.xls'):
        return read_excel(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_path}")
